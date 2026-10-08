#!/usr/bin/env python3
"""Force a .pptx to the ServiceNow standard custom size: 13.33 in x 7.5 in.

Usage: python enforce_slide_size.py deck.pptx [out.pptx] [--no-scale]

If the file was a different size, every shape on every slide is scaled
proportionally so content keeps its relative position (pass --no-scale to skip).
Run this as the LAST step before delivering any PowerPoint file.
"""
import sys
from pptx import Presentation
from pptx.util import Emu

SN_WIDTH = Emu(12188952)   # 13.33 in
SN_HEIGHT = Emu(6858000)   # 7.50 in


def _scale_shapes(shapes, sx, sy):
    for shp in shapes:
        try:
            if shp.left is not None:
                shp.left = int(shp.left * sx)
                shp.top = int(shp.top * sy)
            if shp.width is not None:
                shp.width = int(shp.width * sx)
                shp.height = int(shp.height * sy)
        except (AttributeError, ValueError):
            pass


def enforce(src, dst=None, scale=True):
    prs = Presentation(src)
    ow, oh = prs.slide_width, prs.slide_height
    sx, sy = SN_WIDTH / ow, SN_HEIGHT / oh
    prs.slide_width, prs.slide_height = SN_WIDTH, SN_HEIGHT
    if scale and (abs(sx - 1) > 0.001 or abs(sy - 1) > 0.001):
        for slide in prs.slides:
            _scale_shapes(slide.shapes, sx, sy)
    prs.save(dst or src)
    print(f"{dst or src}: {prs.slide_width / 914400:.2f} x {prs.slide_height / 914400:.2f} in "
          f"(was {ow / 914400:.2f} x {oh / 914400:.2f} in)")


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        sys.exit(__doc__)
    enforce(args[0], args[1] if len(args) > 1 else None, "--no-scale" not in sys.argv)
