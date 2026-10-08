#!/usr/bin/env python3
"""Print every layout in the Corporate Template and its placeholders."""
import os, sys
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "..", "templates", "ServiceNow_Corporate_Template_v2.3.pptx")

prs = Presentation(sys.argv[1] if len(sys.argv) > 1 else TEMPLATE)
for m_i, master in enumerate(prs.slide_masters):
    for layout in master.slide_layouts:
        print(f"\n[{layout.name.strip()}]")
        for ph in sorted(layout.placeholders, key=lambda p: (p.top or 0, p.left or 0)):
            pf = ph.placeholder_format
            print(f"  idx={pf.idx:<3} {str(pf.type):<22} '{ph.name}'  "
                  f"@ {(ph.left or 0)/914400:.2f},{(ph.top or 0)/914400:.2f} in")
