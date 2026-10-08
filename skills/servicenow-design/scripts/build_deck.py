#!/usr/bin/env python3
"""Build an on-brand ServiceNow deck from a JSON outline using the Corporate Template.

Usage: python build_deck.py outline.json out.pptx

Outline format (see templates/deck-outline.example.json):
{
  "slides": [
    {"layout": "Cover Green", "title": "...", "subtitle": "..."},
    {"layout": "Title and Content", "title": "...", "subtitle": "...",
     "body": [["Point one", {"text": "Sub point", "level": 1}]],
     "images": ["path/to/picture.png"], "notes": "Speaker notes"}
  ]
}
- title    -> title / centre-title placeholder
- subtitle -> subtitle placeholder (or first body placeholder above the content)
- body     -> remaining text placeholders in reading order (top-left first);
              each entry is a string or a list of paragraphs (string or {text, level})
- images   -> picture placeholders in reading order
Output is always 13.33 x 7.5 in.
"""
import json, os, sys
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import PP_PLACEHOLDER

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "..", "templates", "ServiceNow_Corporate_Template_v2.3.pptx")
SN_WIDTH, SN_HEIGHT = Emu(12188952), Emu(6858000)

TITLE = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
SUBTITLE = {PP_PLACEHOLDER.SUBTITLE}
TEXT = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT}
PICTURE = {PP_PLACEHOLDER.PICTURE}
CHROME = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.SLIDE_NUMBER}


def remove_all_slides(prs):
    lst = prs.slides._sldIdLst
    for sld in list(lst):
        prs.part.drop_rel(sld.rId)
        lst.remove(sld)


def layouts_by_name(prs):
    out = {}
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            out.setdefault(layout.name.strip().lower(), layout)
    return out


def fill_text(ph, content):
    paras = content if isinstance(content, list) else [content]
    tf = ph.text_frame
    tf.clear()
    for i, p in enumerate(paras):
        text, level = (p.get("text", ""), p.get("level", 0)) if isinstance(p, dict) else (str(p), 0)
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = text
        para.level = level


def build(outline_path, out_path):
    with open(outline_path, encoding="utf-8") as f:
        outline = json.load(f)
    base = os.path.dirname(os.path.abspath(outline_path))
    prs = Presentation(TEMPLATE)
    remove_all_slides(prs)
    layouts = layouts_by_name(prs)

    for n, spec in enumerate(outline["slides"], 1):
        name = spec.get("layout", "Title and Content").strip().lower()
        if name == "do not use":
            sys.exit(f"Slide {n}: layout 'Do not use' is forbidden")
        if name not in layouts:
            sys.exit(f"Slide {n}: unknown layout '{spec.get('layout')}'. Run list_layouts.py")
        slide = prs.slides.add_slide(layouts[name])
        phs = sorted(slide.placeholders, key=lambda p: (p.top or 0, p.left or 0))
        used = set()

        def take(types):
            for ph in phs:
                if ph.placeholder_format.idx not in used and ph.placeholder_format.type in types:
                    used.add(ph.placeholder_format.idx)
                    return ph

        if spec.get("title") is not None and (ph := take(TITLE)):
            fill_text(ph, spec["title"])
        if spec.get("subtitle") is not None and (ph := take(SUBTITLE) or take(TEXT)):
            fill_text(ph, spec["subtitle"])
        for block in spec.get("body", []):
            ph = take(TEXT | SUBTITLE)
            if ph is None:
                print(f"warning: slide {n} has more body blocks than text placeholders", file=sys.stderr)
                break
            fill_text(ph, block)
        for img in spec.get("images", []):
            ph = take(PICTURE)
            if ph is None:
                print(f"warning: slide {n} has more images than picture placeholders", file=sys.stderr)
                break
            ph.insert_picture(os.path.join(base, img))
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]

        # remove unfilled content placeholders so no "Click to add" prompts remain
        for ph in list(slide.placeholders):
            pf = ph.placeholder_format
            if pf.idx not in used and pf.type not in CHROME:
                ph._element.getparent().remove(ph._element)

    prs.slide_width, prs.slide_height = SN_WIDTH, SN_HEIGHT
    prs.save(out_path)
    print(f"{out_path}: {len(prs.slides)} slides, "
          f"{prs.slide_width / 914400:.2f} x {prs.slide_height / 914400:.2f} in")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    build(sys.argv[1], sys.argv[2])
